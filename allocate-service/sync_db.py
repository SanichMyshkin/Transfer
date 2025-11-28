import os
import logging
import sqlite3
import pymssql
from dotenv import load_dotenv
from datetime import datetime, date, time
from openpyxl import Workbook

load_dotenv()

logger = logging.getLogger()
logger.setLevel(logging.INFO)

file_handler = logging.FileHandler("sync.log", encoding="utf-8")
console_handler = logging.StreamHandler()

formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
file_handler.setFormatter(formatter)
console_handler.setFormatter(formatter)

logger.addHandler(file_handler)
logger.addHandler(console_handler)

SQLITE_FILE = os.getenv("SQLITE_FILE")
SQLITE_TABLE = os.getenv("SQLITE_TABLE")
SOURCE_VIEW = os.getenv("SOURCE_VIEW")
BATCH_SIZE = int(os.getenv("BATCH_SIZE", "1000"))


def normalize_value(v):
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    return v


def fetch_mssql_data():
    conn = pymssql.connect(
        server=os.getenv("MSSQL_SERVER"),
        user=os.getenv("MSSQL_USER"),
        password=os.getenv("MSSQL_PASSWORD"),
        database=os.getenv("MSSQL_DB"),
    )
    with conn:
        with conn.cursor() as cursor:
            cursor.execute(f"SELECT * FROM {SOURCE_VIEW}")
            rows = cursor.fetchall()
            columns = [col[0] for col in cursor.description]
            types = [str(col[1]).lower() for col in cursor.description]
            return columns, rows, types


def map_sql_type(t):
    if "int" in t:
        return "INTEGER"
    if "decimal" in t or "numeric" in t or "float" in t or "real" in t:
        return "REAL"
    if "date" in t or "time" in t:
        return "TEXT"
    if "char" in t or "text" in t:
        return "TEXT"
    if "bit" in t:
        return "INTEGER"
    return "TEXT"


def ensure_table_exists(columns, types):
    with sqlite3.connect(SQLITE_FILE) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
            (SQLITE_TABLE,),
        )
        if cur.fetchone():
            logger.info(f"Таблица {SQLITE_TABLE} существует.")
            return

        logger.info(f"Создание таблицы {SQLITE_TABLE}.")
        defs = [f"{n} {map_sql_type(t)}" for n, t in zip(columns, types)]
        create_sql = f"CREATE TABLE {SQLITE_TABLE} ({', '.join(defs)});"
        cur.execute(create_sql)
        conn.commit()
        logger.info(f"Таблица {SQLITE_TABLE} создана.")


def load_into_sqlite(columns, rows):
    with sqlite3.connect(SQLITE_FILE) as conn:
        cur = conn.cursor()
        logger.info(f"Очистка таблицы {SQLITE_TABLE}.")
        cur.execute(f"DELETE FROM {SQLITE_TABLE}")

        placeholders = ", ".join(["?"] * len(columns))
        collist = ", ".join(columns)
        insert_sql = f"INSERT INTO {SQLITE_TABLE} ({collist}) VALUES ({placeholders})"

        batch = []
        count = 0

        for row in rows:
            norm_row = tuple(normalize_value(v) for v in row)
            batch.append(norm_row)
            count += 1

            if len(batch) >= BATCH_SIZE:
                cur.executemany(insert_sql, batch)
                logger.info(f"Вставлено строк: {count}")
                batch = []

        if batch:
            cur.executemany(insert_sql, batch)
            logger.info(f"Вставлено строк: {count}")

        conn.commit()
        logger.info("Загрузка завершена.")


def log_first_20_rows():
    with sqlite3.connect(SQLITE_FILE) as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM {SQLITE_TABLE} LIMIT 20")
        rows = cur.fetchall()
        logger.info("Первые 20 строк таблицы:")
        for r in rows:
            logger.info(str(r))


def export_to_excel(columns):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    with sqlite3.connect(SQLITE_FILE) as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM {SQLITE_TABLE}")
        rows = cur.fetchall()

    ws.append(columns)

    for r in rows:
        ws.append(list(r))

    wb.save("export.xlsx")
    logger.info("Excel файл export.xlsx создан.")


def main():
    logger.info("Старт обработки.")
    columns, rows, types = fetch_mssql_data()
    logger.info(f"Получено строк: {len(rows)}")
    logger.info(f"Колонки: {columns}")

    ensure_table_exists(columns, types)
    load_into_sqlite(columns, rows)
    log_first_20_rows()
    export_to_excel(columns)

    logger.info("Готово.")


if __name__ == "__main__":
    main()
